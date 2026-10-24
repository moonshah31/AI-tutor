from datetime import UTC, datetime
from typing import BinaryIO

from pptx import Presentation

from app.extractors.base import DocumentExtractor
from app.models.extracted_document import DocumentSection


class PPTXExtractor(DocumentExtractor):
    """Extracts structured text from PowerPoint presentations."""

    def extract(
        self,
        file: BinaryIO,
    ) -> list[DocumentSection]:
        presentation = Presentation(file)

        sections: list[DocumentSection] = []

        for slide_number, slide in enumerate(
            presentation.slides,
            start=1,
        ):
            text_parts: list[str] = []

            for shape in slide.shapes:
                if not hasattr(shape, "text"):
                    continue

                text = shape.text.strip()

                if text:
                    text_parts.append(text)

            if not text_parts:
                continue

            sections.append(
                {
                    "section_number": slide_number,
                    "location_label": f"Slide {slide_number}",
                    "text": "\n".join(text_parts),
                    "created_at": datetime.now(UTC),
                }
            )

        return sections